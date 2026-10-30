"""Tests for the internal AI service that run WITHOUT a Gemini key.

They cover the internal-token boundary, the deterministic fallbacks, the
weighted final score, DOCX generation and resume parsing. A real Gemini
(text or Live) session is NOT exercised here.
"""

import io
import os

os.environ["AI_SERVICE_TOKEN"] = "test-token"
os.environ["GEMINI_API_KEY"] = ""

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from starlette.websockets import WebSocketDisconnect  # noqa: E402

from app.main import app  # noqa: E402

client = TestClient(app)
AUTH = {"X-Internal-Token": "test-token"}
RESUME = "Aarav Sharma\nSoftware Engineer\naarav@example.com | +91 98765 43210\n\nSKILLS\nPython, FastAPI, SQL\n\nEXPERIENCE\nIntern | TechNova\n- Built REST APIs with FastAPI"
JD = "Backend engineer with Python, FastAPI, SQL, Docker and Kubernetes experience building REST APIs."


def test_health_is_public_and_reports_gemini_unconfigured():
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["gemini_live_configured"] is False


@pytest.mark.parametrize("path", ["/api/interview/start", "/api/interview/evaluate", "/api/interview/final", "/api/resume/generate", "/api/ai/rewrite-bullet"])
def test_api_rejects_requests_without_internal_token(path):
    assert client.post(path, json={}).status_code == 401
    assert client.post(path, json={}, headers={"X-Internal-Token": "wrong"}).status_code == 401


def test_live_websocket_rejects_missing_token():
    with pytest.raises(WebSocketDisconnect):
        with client.websocket_connect("/ws/monk-live") as ws:
            ws.receive_json()


def test_live_websocket_reports_missing_gemini_key():
    with client.websocket_connect("/ws/monk-live", headers=AUTH) as ws:
        msg = ws.receive_json()
        assert msg["type"] == "error"
        assert "GEMINI_API_KEY" in msg["message"]


def test_interview_start_falls_back_and_says_so():
    res = client.post("/api/interview/start", json={"resume_text": RESUME, "job_description": JD}, headers=AUTH)
    assert res.status_code == 200
    body = res.json()
    assert len(body["questions"]) == 12
    assert body["ai_used"] is False
    assert "GEMINI_API_KEY" in body["fallback_reason"]


def test_evaluate_passes_through_only_measured_presentation():
    res = client.post(
        "/api/interview/evaluate",
        json={"question": "Explain REST", "answer": "REST APIs use HTTP verbs; I built them with FastAPI and SQL.", "resume_text": RESUME, "job_description": JD, "vision_metrics": {"facial": 61.4, "eyeContact": None}},
        headers=AUTH,
    )
    assert res.status_code == 200
    body = res.json()
    assert body["ai_used"] is False
    assert body["presentation"] == {"Facial_Engagement": 61}
    assert 0 <= body["score"] <= 100


def test_final_score_excludes_unmeasured_components_and_renormalizes():
    history = [
        {"round": "Technical", "question": "q1", "answer": "a", "score": 80, "vision": {}},
        {"round": "Technical", "question": "q2", "answer": "b", "score": 60, "vision": {}},
    ]
    body = client.post("/api/interview/final", json={"resume_text": RESUME, "job_description": JD, "history": history}, headers=AUTH).json()
    # Only answer quality was measured -> overall equals its average (70), not 70*0.4=28.
    assert body["overall_score"] == 70
    assert set(body["unavailable_components"]) == {"communication", "facial_expression", "body_language", "response_delivery"}
    assert body["ai_used"] is False


def test_final_score_uses_requested_weights_when_all_measured():
    vision = {"communication": 50, "facial": 50, "posture": 50, "bodyMovement": 50}
    history = [{"round": "Technical", "question": "q", "answer": "a", "score": 100, "vision": vision}]
    body = client.post("/api/interview/final", json={"resume_text": RESUME, "job_description": JD, "history": history}, headers=AUTH).json()
    assert body["overall_score"] == round(100 * 0.40 + 50 * 0.60)
    assert body["unavailable_components"] == []
    assert body["components"]["answer_technical_quality"] == {"weight": 40, "score": 100, "samples": 1}


def test_rewrite_requires_gemini():
    res = client.post("/api/ai/rewrite-bullet", json={"text": "Worked on the website"}, headers=AUTH)
    assert res.status_code == 503


@pytest.mark.parametrize("template", ["jake", "faangpath", "awesome", "rendercv", "deedy", "classic", "modern"])
def test_docx_generation_for_every_template(template):
    res = client.post("/api/resume/generate", json={"resume_text": RESUME, "template": template, "filename": "Test"}, headers=AUTH)
    assert res.status_code == 200
    assert res.content[:2] == b"PK"
    assert "wordprocessingml" in res.headers["content-type"]
    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(res.content)).paragraphs)
    for table in Document(io.BytesIO(res.content)).tables:
        for row in table.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    assert "Aarav Sharma" in text
    assert "Built REST APIs with FastAPI" in text


def test_docx_rejects_unknown_template():
    res = client.post("/api/resume/generate", json={"resume_text": RESUME, "template": "neon"}, headers=AUTH)
    assert res.status_code == 400


def test_parse_txt_and_reject_unsupported_types():
    ok = client.post("/api/resume/parse", files={"file": ("cv.txt", RESUME.encode(), "text/plain")}, headers=AUTH)
    assert ok.status_code == 200
    assert ok.json()["structured"]["personal"]["email"] == "aarav@example.com"
    bad = client.post("/api/resume/parse", files={"file": ("cv.exe", b"MZ...", "application/octet-stream")}, headers=AUTH)
    assert bad.status_code == 415
